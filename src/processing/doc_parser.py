# src/processing/doc_parser.py
"""
Extract text from Microsoft Word (.docx) documents.
Handles paragraphs and table cell text.
"""

from pathlib import Path
from typing import Tuple

def extract_docx_content(docx_path: Path) -> Tuple[str, float]:
    """
    Extract text from a .docx file.
    Returns (text, confidence) where confidence is always 1.0
    because we're reading the actual text, not OCR.
    """
    try:
        from docx import Document
    except ImportError:
        raise ImportError(
            "python-docx is required for Word documents. "
            "Install it with: pip install python-docx"
        )

    doc = Document(str(docx_path))
    paragraphs = []
    # Extract paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)

    # Extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())
            if row_text:
                paragraphs.append(" | ".join(row_text))

    full_text = "\n".join(paragraphs)
    # Word documents have perfect text (no OCR), so confidence = 1.0
    return full_text, 1.0